#!/usr/bin/env python3
"""Build a PowerPoint-365-compatible PPTX from the mentoring-ideation SVGs.

PowerPoint's OOXML format stores an SVG picture as TWO parts:
  1. a raster PNG fallback referenced by <a:blip r:embed="rIdX"/>, and
  2. the vector SVG itself, referenced through the Microsoft extension
     <a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">
       <asvg:svgBlip r:embed="rIdY"/>
     </a:ext>
A blip that points straight at an SVG part with no PNG fallback is rejected
by PowerPoint ("needs repair" / "cannot read"), which is what happened with
the earlier hand-built attempt. python-pptx 1.0.2 alone cannot emit
svgBlip, so this script:

  1. renders each SVG to a 2x-resolution PNG fallback with cairosvg
     (temp files only — never delivered),
  2. builds a standard schema-valid package with python-pptx and places
     each PNG full-bleed on a blank 16:9 slide (13.333in x 7.5in),
  3. post-processes the zip to inject the svgBlip extension into every
     a:blip, add the SVG media parts, their image relationships, and the
     svg default content-type.

PowerPoint 365 then renders the vector SVG (what you see and edit); the
PNG is only a compatibility fallback, exactly as PowerPoint itself saves
SVG pictures.

Output: slides/mentoring-ideation/build/pptx/serambi-mentoring-ideation.pptx
"""

import os
import re
import tempfile
import zipfile

import cairosvg
from defusedxml import ElementTree as SafeET
from pptx import Presentation
from pptx.util import Emu

BASE = "/home/adm_ub24/projects/ibm/serambi-ai"
SVG_DIR = os.path.join(BASE, "slides/mentoring-ideation/svg")
OUT_DIR = os.path.join(BASE, "slides/mentoring-ideation/build/pptx")
OUT_FILE = os.path.join(OUT_DIR, "serambi-mentoring-ideation.pptx")

SLIDE_W_EMU = 12192000  # 13.333 in
SLIDE_H_EMU = 6858000   # 7.5 in
FALLBACK_W = 2560       # 2x render
FALLBACK_H = 1440

ASVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
IMG_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"


def read_file(path):
    try:
        with open(path, "rb") as f:
            return f.read()
    except (FileNotFoundError, OSError) as exc:
        raise SystemExit("ERROR: cannot read " + path + ": " + str(exc)) from exc


def render_png_fallback(svg_path, png_path):
    try:
        cairosvg.svg2png(
            url=svg_path,
            write_to=png_path,
            output_width=FALLBACK_W,
            output_height=FALLBACK_H,
            background_color="#FFFFFF",
        )
    except Exception as exc:
        raise SystemExit("ERROR: cairosvg failed on " + svg_path + ": " + str(exc)) from exc


def build_skeleton(png_paths, tmp_pptx):
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)
    blank_layout = prs.slide_layouts[6]  # "Blank" in the default template
    for png in png_paths:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            png, Emu(0), Emu(0), width=Emu(SLIDE_W_EMU), height=Emu(SLIDE_H_EMU)
        )
    prs.save(tmp_pptx)


def postprocess(tmp_pptx, svg_paths, out_file):
    with zipfile.ZipFile(tmp_pptx, "r") as zin:
        entries = {name: zin.read(name) for name in zin.namelist()}

    for i, svg_path in enumerate(svg_paths, start=1):
        slide_part = "ppt/slides/slide" + str(i) + ".xml"
        rels_part = "ppt/slides/_rels/slide" + str(i) + ".xml.rels"

        slide_xml = entries[slide_part].decode("utf-8")
        m = re.search(r'<a:blip r:embed="(rId\d+)"[^>]*/>', slide_xml)
        if not m:
            raise SystemExit("ERROR: no a:blip found in " + slide_part)
        png_rid = m.group(1)

        rels_xml = entries[rels_part].decode("utf-8")
        used = [int(x) for x in re.findall(r'Id="rId(\d+)"', rels_xml)]
        svg_rid = "rId" + str(max(used) + 1)

        new_blip = (
            '<a:blip r:embed="' + png_rid + '">'
            '<a:extLst><a:ext uri="' + ASVG_EXT_URI + '">'
            '<asvg:svgBlip xmlns:asvg="' + ASVG_NS + '" r:embed="' + svg_rid + '"/>'
            "</a:ext></a:extLst></a:blip>"
        )
        slide_xml = slide_xml.replace(m.group(0), new_blip)
        entries[slide_part] = slide_xml.encode("utf-8")

        rels_xml = rels_xml.replace(
            "</Relationships>",
            '  <Relationship Id="' + svg_rid + '" Type="' + IMG_REL_TYPE
            + '" Target="../media/image' + str(i) + '.svg"/>\n</Relationships>',
        )
        entries[rels_part] = rels_xml.encode("utf-8")

        entries["ppt/media/image" + str(i) + ".svg"] = read_file(svg_path)

    ct = entries["[Content_Types].xml"].decode("utf-8")
    if 'Extension="svg"' not in ct:
        ct = ct.replace(
            "</Types>",
            '  <Default Extension="svg" ContentType="image/svg+xml"/>\n</Types>',
        )
        entries["[Content_Types].xml"] = ct.encode("utf-8")

    with zipfile.ZipFile(out_file, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries.items():
            zout.writestr(name, data)


def validate(out_file, num_slides):
    z = zipfile.ZipFile(out_file)
    names = set(z.namelist())
    errors = []

    for name in z.namelist():
        if name.endswith((".xml", ".rels")):
            try:
                SafeET.fromstring(z.read(name))
            except Exception as exc:
                errors.append("parse " + name + ": " + str(exc))

    for name in [n for n in z.namelist() if n.endswith(".rels")]:
        base = os.path.dirname(os.path.dirname(name))
        for rel in SafeET.fromstring(z.read(name)):
            if rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target")
            if not target:
                errors.append(name + ": relationship without Target")
                continue
            resolved = os.path.normpath(os.path.join(base, target)).replace(os.sep, "/")
            if resolved not in names:
                errors.append(name + ": unresolved target " + target + " -> " + resolved)

    for i in range(1, num_slides + 1):
        xml = z.read("ppt/slides/slide" + str(i) + ".xml").decode("utf-8")
        if "svgBlip" not in xml:
            errors.append("slide" + str(i) + ": svgBlip extension missing")
        m = re.findall(r'<asvg:svgBlip[^>]*r:embed="(rId\d+)"', xml)
        rels = z.read("ppt/slides/_rels/slide" + str(i) + ".xml.rels").decode("utf-8")
        for rid in m:
            if 'Id="' + rid + '"' not in rels:
                errors.append("slide" + str(i) + ": " + rid + " not in slide rels")
        if "ppt/media/image" + str(i) + ".svg" not in names:
            errors.append("slide" + str(i) + ": SVG media part missing")

    return errors


def main():
    try:
        svg_files = sorted(f for f in os.listdir(SVG_DIR) if f.endswith(".svg"))
    except OSError as exc:
        raise SystemExit("ERROR: cannot list " + SVG_DIR + ": " + str(exc)) from exc
    if not svg_files:
        raise SystemExit("ERROR: no SVG files found in " + SVG_DIR)
    svg_paths = [os.path.join(SVG_DIR, f) for f in svg_files]
    num_slides = len(svg_paths)
    print("Found " + str(num_slides) + " SVGs: " + ", ".join(svg_files))

    try:
        os.makedirs(OUT_DIR, exist_ok=True)
    except OSError as exc:
        raise SystemExit("ERROR: cannot create " + OUT_DIR + ": " + str(exc)) from exc
    if os.path.exists(OUT_FILE):
        try:
            os.remove(OUT_FILE)
        except OSError as exc:
            raise SystemExit("ERROR: cannot remove " + OUT_FILE + ": " + str(exc)) from exc

    with tempfile.TemporaryDirectory() as tmp:
        png_paths = []
        for i, svg in enumerate(svg_paths, start=1):
            png = os.path.join(tmp, "slide" + str(i) + ".png")
            render_png_fallback(svg, png)
            png_paths.append(png)
        print("PNG fallbacks rendered (" + str(FALLBACK_W) + "x" + str(FALLBACK_H) + ")")

        tmp_pptx = os.path.join(tmp, "skeleton.pptx")
        build_skeleton(png_paths, tmp_pptx)
        print("python-pptx skeleton built")

        postprocess(tmp_pptx, svg_paths, OUT_FILE)

    errors = validate(OUT_FILE, num_slides)
    if errors:
        for e in errors:
            print("FAIL: " + e)
        raise SystemExit("ERROR: validation failed")

    print("Validation: OK (" + str(num_slides) + " slides, svgBlip on every slide, all rels resolve)")
    print("PPTX written: " + OUT_FILE)
    print("Size: " + str(os.path.getsize(OUT_FILE)) + " bytes")


if __name__ == "__main__":
    main()
